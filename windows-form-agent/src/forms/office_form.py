"""Remplissage de modeles Word/Excel.

Convention pour les modeles Word : placeholders au format {{nom_du_champ}}
dans le texte du document (paragraphes ou tableaux).

Convention pour les modeles Excel : une feuille "mapping" (ou la premiere
feuille) ou chaque ligne est "nom_du_champ | reference_cellule" (ex:
"nom_client | B2"), OU un mapping fourni directement en config.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def fill_docx_template(
    template_path: str | Path, output_path: str | Path, values: dict[str, str]
) -> list[str]:
    """Remplace les {{placeholders}} dans un .docx. Renvoie les placeholders non remplaces."""
    doc = Document(str(template_path))

    def replace_in_paragraph(paragraph) -> None:
        for key, value in values.items():
            placeholder = "{{" + key + "}}"
            if placeholder in paragraph.text:
                for run in paragraph.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, value)
                if placeholder in paragraph.text:
                    full_text = "".join(r.text for r in paragraph.runs)
                    full_text = full_text.replace(placeholder, value)
                    for run in paragraph.runs:
                        run.text = ""
                    if paragraph.runs:
                        paragraph.runs[0].text = full_text

    for paragraph in doc.paragraphs:
        replace_in_paragraph(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_in_paragraph(paragraph)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    remaining = extract_remaining_placeholders(output_path)
    return remaining


def extract_remaining_placeholders(docx_path: str | Path) -> list[str]:
    doc = Document(str(docx_path))
    remaining = set()
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)

    for text in texts:
        start = 0
        while True:
            start = text.find("{{", start)
            if start == -1:
                break
            end = text.find("}}", start)
            if end == -1:
                break
            remaining.add(text[start : end + 2])
            start = end + 2

    return sorted(remaining)


def fill_xlsx_template(
    template_path: str | Path,
    output_path: str | Path,
    values: dict[str, str],
    cell_map: dict[str, str],
    sheet_name: str | None = None,
) -> list[str]:
    """Ecrit `values` dans les cellules indiquees par `cell_map` {champ: 'B2'}.

    Renvoie la liste des champs sans cellule associee dans cell_map.
    """
    wb = load_workbook(str(template_path))
    sheet = wb[sheet_name] if sheet_name else wb.active

    missing_mapping = [name for name in values if name not in cell_map]

    for name, cell_ref in cell_map.items():
        if name in values:
            sheet[cell_ref] = values[name]

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))

    return missing_mapping
