"""Extraction de texte depuis des documents Word (.docx) et Excel (.xlsx)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook


def extract_text_from_docx(path: str | Path) -> str:
    doc = Document(str(path))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_text_from_xlsx(path: str | Path) -> str:
    wb = load_workbook(str(path), data_only=True, read_only=True)
    parts: list[str] = []

    for sheet in wb.worksheets:
        parts.append(f"# Feuille: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                parts.append(" | ".join(values))

    return "\n".join(parts).strip()
